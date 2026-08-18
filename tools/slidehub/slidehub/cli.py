"""slidehub — P0 feasibility spike.

The question this tool exists to answer: can a page be lifted out of one deck,
dropped into another, and still look like itself? Everything else in the platform
plan depends on the answer, so it gets measured before any UI is built.

    python3 -m slidehub spike                      # full end-to-end verdict
    python3 -m slidehub ingest <deck.pptx> ...     # split + render + dedup
    python3 -m slidehub assemble --pages A:1,C:2   # build a deck
    python3 -m slidehub roundtrip <deck.pptx>      # read origin markers back
"""
from __future__ import annotations

import argparse
import json
import sys
import time
from pathlib import Path

from pptx import Presentation

from . import catalog, library, manifest, marker, pdfdoc, reimport, simulate, verify
from .assemble import assemble
from .fingerprint import COLOR_WARN, FIDELITY_WARN

HERE = Path(__file__).resolve().parent.parent


def _rule(title: str = "") -> None:
    print("\n" + (" " + title + " ").center(74, "=") if title else "=" * 74)


def _fmt_report(rep) -> str:
    return ("pages=%d  exact=%d  layout(mean=%.2f worst=%d)  "
            "colour(mean=%.2f worst=%.2f)  failing=%d"
            % (rep.pages, rep.exact, rep.mean, rep.worst,
               rep.mean_color, rep.worst_color, len(rep.failures)))


# --------------------------------------------------------------------------- #
def cmd_ingest(args) -> int:
    decks = sorted(Path(p) for p in args.decks)
    if not decks:
        print("no decks given", file=sys.stderr)
        return 2
    print("ingesting %d deck(s) -> %s" % (len(decks), args.out))
    index = library.ingest(decks, Path(args.out), dpi=args.dpi)

    total = len(index["pages"])
    modules = len(index["modules"])
    _rule("ingest")
    for d in index["decks"]:
        print("  %s  %-28s %2d pages" % (d["id"], d["name"], d["slides"]))
    print("\n  %d pages  ->  %d distinct modules  (%d duplicate pages collapsed)"
          % (total, modules, total - modules))
    dupes = [m for m in index["modules"] if m["version_count"] > 1]
    if dupes:
        print("\n  modules with more than one version:")
        for m in dupes:
            print("    %s  x%d  %s  [%s]"
                  % (m["module_id"], m["version_count"], m["title"],
                     ", ".join(m["members"])))
    return 0


def cmd_assemble(args) -> int:
    index = library.load(Path(args.library))
    pages = library.page_map(index)
    uids = [u.strip() for u in args.pages.split(",") if u.strip()]
    missing = [u for u in uids if u not in pages]
    if missing:
        print("unknown page(s): %s" % ", ".join(missing), file=sys.stderr)
        return 2

    items = [{"pptx": pages[u]["pptx"], "module_id": pages[u]["module_id"],
              "version_id": u} for u in uids]
    rep = assemble(items, Path(args.out), template=args.template,
                   export_id=args.export_id)
    print("assembled %d pages -> %s" % (rep.slides, args.out))
    print("  parts cloned=%d  relationships rewritten=%d  masters carried=%d"
          % (rep.parts_cloned, rep.rels_rewritten, rep.masters_added))
    for w in rep.size_mismatches:
        print("  ! slide-size mismatch: %s" % w)
    for w in rep.warnings:
        print("  ! %s" % w)
    return 0


def cmd_roundtrip(args) -> int:
    prs = Presentation(args.deck)
    print("reading origin markers from %s" % args.deck)
    found = 0
    for i, slide in enumerate(prs.slides, start=1):
        info = marker.read(slide)
        if info:
            found += 1
            print("  page %2d  module=%s  version=%s  export=%s"
                  % (i, info["module_id"], info["version_id"], info["export_id"]))
        else:
            print("  page %2d  <no marker — would fall back to fingerprint match>" % i)
    print("\n  %d/%d pages identified" % (found, len(prs.slides)))
    return 0 if found == len(prs.slides) else 1


def _load_catalog(path):
    return json.loads((Path(path) / "index.json").read_text(encoding="utf-8"))


def cmd_catalog(args) -> int:
    print("building catalogue from %d source(s)" % len(args.sources))
    index = catalog.build([Path(s) for s in args.sources], Path(args.out),
                          thumb_width=args.thumb_width)
    multi = [m for m in index["modules"] if m["version_count"] > 1]
    flagged = [p for p in index["pages"] if p["review"]["status"] != "pending"]
    weak = [p for p in index["pages"] if len(p["review"]["weak_fields"]) >= 3]
    _rule("catalogue")
    print("  %d pages  ->  %d groups  ->  %d modules" %
          (len(index["pages"]), len(index["groups"]), len(index["modules"])))
    print("  %d modules carry more than one version (the same case in several decks)"
          % len(multi))
    print("\n  review queue:")
    print("    %2d pages flagged as attached to the wrong case" % len(flagged))
    print("    %2d pages where three or more tag fields are low-confidence" % len(weak))
    return 0


def _matches(page, args, terms) -> bool:
    tags = page["tags"]
    facets = (("year", [tags["year"]]), ("region", tags["regions"]),
              ("industry", tags["industries"]), ("service", tags["services"]),
              ("event", tags["events"]), ("client", tags["clients"]),
              ("type", [tags["page_type"]]))
    for name, values in facets:
        wanted = getattr(args, name, None)
        if wanted and not any(wanted in (v or "") for v in values):
            return False
    if terms:
        blob = "%s %s %s" % (page["title"], page["text"], " ".join(tags["clients"]))
        if not all(t.lower() in blob.lower() for t in terms):
            return False
    return True


def cmd_search(args) -> int:
    index = _load_catalog(args.library)
    hits = [p for p in index["pages"] if _matches(p, args, args.query)]
    print("%d / %d pages match" % (len(hits), len(index["pages"])))
    for page in hits[:args.limit]:
        tags = page["tags"]
        print("  %-6s %-38s %-5s %-8s %-14s %s"
              % (page["uid"], page["title"][:38], tags["year"] or "—",
                 "/".join(tags["regions"][:1]) or "—",
                 "/".join(tags["industries"][:1]) or "—",
                 "/".join(tags["clients"][:2]) or "—"))
    if len(hits) > args.limit:
        print("  ... %d more" % (len(hits) - args.limit))
    return 0


def cmd_compose(args) -> int:
    index = _load_catalog(args.library)
    by_uid = {p["uid"]: p for p in index["pages"]}
    by_doc = {d["id"]: d["file"] for d in index["docs"]}
    groups = {g["group_id"]: g for g in index["groups"]}
    modules = {m["module_id"]: m for m in index["modules"]}

    uids: list[str] = []
    for token in (t.strip() for t in args.pages.split(",") if t.strip()):
        if token in modules:
            uids.extend(groups[modules[token]["default_version"]]["members"])
        elif token in groups:
            uids.extend(groups[token]["members"])
        elif token in by_uid:
            uids.append(token)
        else:
            print("unknown page / group / module: %s" % token, file=sys.stderr)
            return 2

    refs = [(by_doc[by_uid[u]["doc"]], by_uid[u]["index"]) for u in uids]
    written = pdfdoc.compose(refs, Path(args.out),
                             title=args.title or Path(args.out).stem)

    # "最后一次使用" only means anything if composing writes it back. Recording
    # it here is also what makes the reverse question answerable later: which
    # pages actually go out to clients, and which have never left the library.
    stamp = time.strftime("%Y-%m-%d %H:%M:%S")
    for uid in set(uids):
        by_uid[uid]["last_used"] = stamp
        by_uid[uid]["use_count"] = by_uid[uid].get("use_count", 0) + 1
    index.setdefault("exports", []).append({
        "at": stamp, "out": str(args.out),
        "title": args.title or Path(args.out).stem, "pages": uids,
    })
    (Path(args.library) / "index.json").write_text(
        json.dumps(index, ensure_ascii=False, indent=1), encoding="utf-8")

    print("composed %d pages -> %s" % (written, args.out))
    for n, uid in enumerate(uids, start=1):
        print("  %2d. %-6s %s" % (n, uid, by_uid[uid]["title"][:56]))
    return 0


def _print_verdicts(verdicts) -> dict:
    counts: dict[str, int] = {}
    print("  %-5s %-11s %-9s %-8s %-7s %s"
          % ("page", "status", "module", "method", "textsim", "suggested action"))
    for v in verdicts:
        counts[v.status] = counts.get(v.status, 0) + 1
        pos = "-" if v.position == 0 else str(v.position)
        print("  %-5s %-11s %-9s %-8s %-7s %s"
              % (pos, v.status, v.module_id or "-", v.match_method,
                 "%.2f" % v.text_similarity if v.text_similarity else "-",
                 v.suggestion))
        if v.note:
            print("        note: %s" % v.note)
        if v.candidates:
            print("        candidates: %s" % ", ".join(v.candidates))
    return counts


def cmd_reimport(args) -> int:
    index = library.load(Path(args.library))
    verdicts = reimport.analyse(Path(args.deck), index,
                                Path(args.work), dpi=args.dpi)
    print("re-import analysis of %s" % args.deck)
    counts = _print_verdicts(verdicts)
    print("\n  " + "  ".join("%s=%d" % kv for kv in sorted(counts.items())))
    return 0


def cmd_spike(args) -> int:
    out = Path(args.out)
    lib = out / "library"
    decks = sorted(Path(args.decks).glob("*.pptx"))
    if not decks:
        print("no decks in %s — run samples/make_samples.py first" % args.decks,
              file=sys.stderr)
        return 2

    _rule("P0 · step 1 — ingest")
    index = library.ingest(decks, lib, dpi=args.dpi)
    for d in index["decks"]:
        print("  %s  %-28s %2d pages" % (d["id"], d["name"], d["slides"]))
    print("  %d pages -> %d modules" % (len(index["pages"]), len(index["modules"])))

    # Interleave decks so consecutive pages come from different themes — the
    # hardest case for assembly, and the realistic one for a client deck.
    by_deck: dict[str, list] = {}
    for p in index["pages"]:
        by_deck.setdefault(p["deck_id"], []).append(p)
    picked = []
    row = 0
    while len(picked) < args.pages:
        added = False
        for deck_id in sorted(by_deck):
            if row < len(by_deck[deck_id]) and len(picked) < args.pages:
                picked.append(by_deck[deck_id][row])
                added = True
        if not added:
            break
        row += 1

    _rule("P0 · step 2 — cross-deck selection")
    for i, p in enumerate(picked, start=1):
        print("  %2d. %s  %-11s %s" % (i, p["uid"], "[%s]" % p["deck_id"], p["title"]))

    references = [(p["uid"], Path(p["png"])) for p in picked]
    items = [{"pptx": p["pptx"], "module_id": p["module_id"], "version_id": p["uid"]}
             for p in picked]

    _rule("P0 · step 3 — assemble")
    deck_out = out / "assembled.pptx"
    rep = assemble(items, deck_out, template=args.template, export_id="spike001")
    print("  -> %s" % deck_out)
    print("     parts cloned=%d  rels rewritten=%d  masters carried=%d"
          % (rep.parts_cloned, rep.rels_rewritten, rep.masters_added))
    for w in rep.warnings[:5]:
        print("     ! %s" % w)

    _rule("P0 · step 4 — fidelity check")
    fid = verify.check(deck_out, references, out / "render_assembled",
                       threshold=args.threshold, dpi=args.dpi)
    print("  %s" % _fmt_report(fid))
    for c in fid.checks:
        flag = "ok   " if c.ok else "DRIFT"
        note = "" if c.ok else "  <- %s changed" % c.reason
        print("    %s page %2d  <- %-6s layout=%2d  colour=%5.2f%s"
              % (flag, c.position, c.source_uid, c.distance, c.color_delta, note))

    _rule("P0 · step 5 — round-trip")
    exported = deck_out
    prs = Presentation(str(exported))
    recovered = sum(1 for sl in prs.slides if marker.read(sl))
    has_manifest = manifest.read(prs) is not None
    print("  export carries %d/%d page markers, deck manifest: %s"
          % (recovered, len(prs.slides), "present" if has_manifest else "MISSING"))

    returned = out / "returned_from_client.pptx"
    print("\n  simulating what happens after the meeting:")
    for line in simulate.client_meeting(exported, returned):
        print("    - %s" % line)

    print("\n  analysing the returned deck:")
    verdicts = reimport.analyse(returned, index, out / "render_returned", dpi=args.dpi)
    counts = _print_verdicts(verdicts)

    _rule("VERDICT")
    print("  fidelity   : %d/%d pages bit-identical to source  (worst layout=%d, colour=%.2f)"
          % (fid.exact, fid.pages, fid.worst, fid.worst_color))
    print("  markers    : %d/%d recovered, manifest %s"
          % (recovered, len(prs.slides), "present" if has_manifest else "MISSING"))
    print("  round-trip : %s" % "  ".join("%s=%d" % kv for kv in sorted(counts.items())))
    print("\n  layout threshold = %d bits of a 64-bit perceptual hash" % args.threshold)
    print("  colour threshold = %.1f mean channel delta (0-255)" % COLOR_WARN)
    print("\n  Open the assembled deck in PowerPoint and compare against the sources.")
    print("  The numbers narrow down where to look; your eyes make the call.")
    print("    %s" % deck_out)

    expected = {"unchanged": 7, "modified": 1, "unmatched": 1,
                "new": 1, "deleted": 1}
    round_trip_ok = counts == expected
    if not round_trip_ok:
        print("\n  ! round-trip classification differs from the simulated edits")
        print("    expected %s" % expected)
        print("    got      %s" % counts)
    ok = fid.passed and recovered == len(prs.slides) and has_manifest and round_trip_ok
    return 0 if ok else 1


# --------------------------------------------------------------------------- #
def main(argv=None) -> int:
    ap = argparse.ArgumentParser(prog="slidehub", description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)

    p = sub.add_parser("ingest", help="split decks into pages, render, deduplicate")
    p.add_argument("decks", nargs="+")
    p.add_argument("--out", default="build/library")
    p.add_argument("--dpi", type=int, default=110)
    p.set_defaults(func=cmd_ingest)

    p = sub.add_parser("assemble", help="build a deck from library pages")
    p.add_argument("--library", default="build/library")
    p.add_argument("--pages", required=True, help="comma-separated uids, e.g. A:1,C:2,B:3")
    p.add_argument("--out", default="build/assembled.pptx")
    p.add_argument("--template", default=None)
    p.add_argument("--export-id", default="exp0001")
    p.set_defaults(func=cmd_assemble)

    p = sub.add_parser("roundtrip", help="read origin markers back out of a deck")
    p.add_argument("deck")
    p.set_defaults(func=cmd_roundtrip)

    p = sub.add_parser("reimport", help="analyse a deck returned from a client")
    p.add_argument("--library", default="build/library")
    p.add_argument("--deck", required=True)
    p.add_argument("--work", default="build/reimport")
    p.add_argument("--dpi", type=int, default=110)
    p.set_defaults(func=cmd_reimport)

    p = sub.add_parser("catalog", help="build the searchable PDF catalogue")
    p.add_argument("sources", nargs="+", help="source PDF files")
    p.add_argument("--out", default="build/lib")
    p.add_argument("--thumb-width", type=int, default=480)
    p.set_defaults(func=cmd_catalog)

    p = sub.add_parser("search", help="search the catalogue")
    p.add_argument("query", nargs="*", default=[])
    p.add_argument("--library", default="build/lib")
    for facet in ("year", "region", "industry", "service", "event", "client", "type"):
        p.add_argument("--%s" % facet, default=None)
    p.add_argument("--limit", type=int, default=30)
    p.set_defaults(func=cmd_search)

    p = sub.add_parser("compose", help="build a new PDF from catalogue pages")
    p.add_argument("--library", default="build/lib")
    p.add_argument("--pages", required=True,
                   help="comma-separated uids or module ids, e.g. A:6,B:25,m046")
    p.add_argument("--out", default="build/composed.pdf")
    p.add_argument("--title", default="")
    p.set_defaults(func=cmd_compose)

    p = sub.add_parser("spike", help="run the full P0 feasibility check")
    p.add_argument("--decks", default=str(HERE / "samples" / "decks"))
    p.add_argument("--out", default="build/spike")
    p.add_argument("--template", default=None)
    p.add_argument("--pages", type=int, default=10)
    p.add_argument("--dpi", type=int, default=110)
    p.add_argument("--threshold", type=int, default=FIDELITY_WARN)
    p.set_defaults(func=cmd_spike)

    args = ap.parse_args(argv)
    return args.func(args)


if __name__ == "__main__":
    raise SystemExit(main())

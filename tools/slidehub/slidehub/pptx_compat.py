"""Workarounds for python-pptx behaviour this product trips over.

`PresentationPart.add_slide` derives the new slide's package path from
`len(sldIdLst) + 1` — its docstring says "next available partname", but the
implementation only holds while no slide has ever been removed. Delete one slide
from a ten-slide deck and the next added slide claims `slide10.xml`, which is
already occupied. The package then carries two entries under one name and
PowerPoint reports the file as unreadable.

That matters here more than in most projects: delete-then-add is the normal
shape of every builder edit and every returned deck. So slide creation goes
through this module, which allocates against what the package actually contains.
"""
from __future__ import annotations

from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.packuri import PackURI
from pptx.parts.slide import SlidePart

SLIDE_PARTNAME_TMPL = "/ppt/slides/slide%d.xml"


def free_partname(package, tmpl: str = SLIDE_PARTNAME_TMPL) -> PackURI:
    taken = {str(part.partname) for part in package.iter_parts()}
    n = 1
    while tmpl % n in taken:
        n += 1
    return PackURI(tmpl % n)


def add_slide(prs, slide_layout, clone_placeholders: bool = True):
    """Append a slide, allocating a package path that is genuinely unused."""
    package = prs.part.package
    slide_part = SlidePart.new(free_partname(package), package, slide_layout.part)
    rId = prs.part.relate_to(slide_part, RT.SLIDE)
    prs.slides._sldIdLst.add_sldId(rId)
    slide = slide_part.slide
    if clone_placeholders:
        slide.shapes.clone_layout_placeholders(slide_layout)
    return slide
